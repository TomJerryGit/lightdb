# -*- coding: utf-8 -*-
"""导入导出对话框（P1-4：分块 + 后台线程 + 进度 + 取消）

语义约定（**显式写清楚，避免误解**）：
- 导入按 `DEFAULT_BATCH_SIZE=500` 行/批执行；**取消后停止后续批次，已提交的
  批次不回滚**（每批由 `execute_many` 独立提交，autocommit 语义），结果里
  `cancelled=True` 并给出“已导入 N 行”。
- 批失败时逐行重试定位失败行，收集 `失败行号 / 错误 / 值`，可导出 CSV 报告。
- 导出（CSV / XLSX）在后台线程分块写文件；**取消或失败会删除不完整文件**。

worker 只做 DB/IO，不触碰 QWidget；进度/完成通过 Qt 信号回主线程。
"""
import csv
import json
import os
import threading
import time

from PySide6.QtCore import QObject, QRunnable, QThreadPool, Qt, Signal
from PySide6.QtWidgets import (
    QDialog, QVBoxLayout, QHBoxLayout, QLabel, QPushButton, QLineEdit,
    QFileDialog, QMessageBox, QProgressDialog, QDialogButtonBox,
    QCheckBox, QTableWidget, QTableWidgetItem, QComboBox,
)

DEFAULT_BATCH_SIZE = 500

__all__ = [
    'DEFAULT_BATCH_SIZE', 'ImportCSVDialog', 'ExportCSVDialog',
    'ImportWorker', 'ExportWorker', 'TransferSignals',
    'chunk_rows', 'normalize_import_params', 'execute_chunked_import',
    'write_failure_report', 'write_export_chunked', 'ExportCancelled',
]


# ==================== 纯逻辑（可在 worker/单测直接调用） ====================

def chunk_rows(rows, batch_size=DEFAULT_BATCH_SIZE):
    """把行序列按批切片（生成器）。"""
    size = max(1, int(batch_size or DEFAULT_BATCH_SIZE))
    rows = list(rows)
    for i in range(0, len(rows), size):
        yield rows[i:i + size]


def normalize_import_params(rows, n_cols):
    """CSV 行 → 参数元组：空字符串视作 NULL，列数不足补 None，多余列忽略。"""
    params = []
    for row in rows:
        vals = []
        for i in range(n_cols):
            v = row[i] if i < len(row) else None
            vals.append(None if v == '' else v)
        params.append(tuple(vals))
    return params


def execute_chunked_import(conn, sql, params_list, batch_size=DEFAULT_BATCH_SIZE,
                           progress_cb=None, cancel_cb=None) -> dict:
    """分块执行 INSERT（核心逻辑；worker 线程调用，单测可直接调用）。

    返回 dict:
      success  已成功行数
      failures [{'row': 行号(1-based), 'error': 错误文本, 'values': [...]}]
      processed 已处理行数（含失败）
      cancelled 是否被取消
      total     总行数
    """
    params_list = list(params_list)
    total = len(params_list)
    success = 0
    processed = 0
    failures = []
    cancelled = False

    for start, batch in zip(range(0, total, max(1, int(batch_size))),
                            chunk_rows(params_list, batch_size)):
        if cancel_cb is not None and cancel_cb():
            cancelled = True
            break
        err = None
        try:
            _n, err = conn.execute_many(sql, batch)
        except Exception as exc:                      # noqa: BLE001
            err = str(exc)
        if err:
            # 批失败：逐行重试，定位具体失败行（成功行照常计入）
            for off, params in enumerate(batch):
                if cancel_cb is not None and cancel_cb():
                    cancelled = True
                    break
                rerr = None
                try:
                    _n, rerr = conn.execute(sql, params)
                except Exception as exc:              # noqa: BLE001
                    rerr = str(exc)
                if rerr:
                    failures.append({'row': start + off + 1,
                                     'error': str(rerr),
                                     'values': list(params)})
                else:
                    success += 1
                processed += 1
                if progress_cb is not None:
                    progress_cb(min(processed, total), total)
            if cancelled:
                break
            continue
        success += len(batch)
        processed += len(batch)
        if progress_cb is not None:
            progress_cb(min(processed, total), total)

    if cancel_cb is not None and cancel_cb():
        cancelled = True
    return {'success': success, 'failures': failures,
            'processed': processed, 'cancelled': cancelled, 'total': total}


def write_failure_report(failures, filepath) -> int:
    """把失败行写成 CSV（含 行号 / 错误原因 / 值）；返回条数。"""
    failures = list(failures or [])
    with open(filepath, 'w', encoding='utf-8-sig', newline='') as f:
        writer = csv.writer(f)
        writer.writerow(['行号', '错误', '值'])
        for item in failures:
            values = item.get('values')
            try:
                values_text = json.dumps(values, ensure_ascii=False,
                                         default=str)
            except Exception:
                values_text = str(values)
            writer.writerow([item.get('row', ''), item.get('error', ''),
                             values_text])
    return len(failures)


class ExportCancelled(Exception):
    """导出被取消（内部用；文件已被删除）。"""


def write_export_chunked(filepath, columns, rows, fmt='csv',
                         with_header=True, null_as_empty=True,
                         batch_size=DEFAULT_BATCH_SIZE,
                         progress_cb=None, cancel_cb=None,
                         chunk_delay=0.0):
    """分块写导出文件（CSV / XLSX）。

    - 分块遍历（每批 `batch_size` 行），进度回调 `(done, total)`；
    - `cancel_cb()` 为真 → 抛 `ExportCancelled` 并**删除不完整文件**；
    - 其它异常同样删除不完整文件后重抛；
    返回写入行数。
    """
    columns = list(columns or [])
    rows = list(rows or [])
    total = len(rows)
    fmt = (fmt or 'csv').lower()
    written = 0
    if fmt == 'xlsx':
        # 注：调用方已经持有 rows 列表；这里用普通 Workbook 按批 append，
        # 取消时**不 save**，因此不会产生不完整文件（也无法留下半截 xlsx）。
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = 'Sheet1'
        if with_header:
            ws.append([str(c) for c in columns])
        null_as = '' if null_as_empty else 'NULL'
        try:
            for batch in chunk_rows(rows, batch_size):
                if cancel_cb is not None and cancel_cb():
                    raise ExportCancelled('cancelled')
                for row in batch:
                    ws.append([null_as if v is None else v for v in row])
                written += len(batch)
                if progress_cb is not None:
                    progress_cb(written, total)
                if chunk_delay:
                    time.sleep(chunk_delay)
            wb.save(filepath)
            return written
        except ExportCancelled:
            _safe_remove(filepath)
            raise
        except Exception:
            _safe_remove(filepath)
            raise
        finally:
            try:
                wb.close()
            except Exception:
                pass

    null_as = '' if null_as_empty else 'NULL'
    try:
        with open(filepath, 'w', encoding='utf-8-sig', newline='') as f:
            writer = csv.writer(f)
            if with_header:
                writer.writerow(columns)
            for batch in chunk_rows(rows, batch_size):
                if cancel_cb is not None and cancel_cb():
                    raise ExportCancelled('cancelled')
                for row in batch:
                    writer.writerow([null_as if v is None else v for v in row])
                written += len(batch)
                if progress_cb is not None:
                    progress_cb(written, total)
                if chunk_delay:
                    time.sleep(chunk_delay)
        return written
    except ExportCancelled:
        _safe_remove(filepath)
        raise
    except Exception:
        _safe_remove(filepath)
        raise


def _safe_remove(path):
    try:
        if path and os.path.exists(path):
            os.remove(path)
    except OSError:
        pass


# ==================== 后台 worker ====================

class TransferSignals(QObject):
    progress = Signal(int, int)      # done, total
    finished = Signal(object)        # result dict
    failed = Signal(str)


class ImportWorker(QRunnable):
    """后台分块导入 worker。"""

    def __init__(self, conn, sql, params_list, batch_size=DEFAULT_BATCH_SIZE):
        super().__init__()
        self.setAutoDelete(False)
        self.conn = conn
        self.sql = sql
        self.params_list = list(params_list)
        self.batch_size = int(batch_size or DEFAULT_BATCH_SIZE)
        self.signals = TransferSignals()
        self._cancel_event = threading.Event()

    def cancel(self):
        self._cancel_event.set()

    @property
    def cancel_requested(self) -> bool:
        return self._cancel_event.is_set()

    def run(self):
        try:
            result = execute_chunked_import(
                self.conn, self.sql, self.params_list,
                batch_size=self.batch_size,
                progress_cb=lambda done, total:
                    self.signals.progress.emit(int(done), int(total)),
                cancel_cb=lambda: self._cancel_event.is_set())
            self.signals.finished.emit(result)
        except BaseException as exc:                  # noqa: BLE001
            self.signals.failed.emit(str(exc))


class ExportWorker(QRunnable):
    """后台分块导出 worker（CSV/XLSX，取消删除不完整文件）。"""

    def __init__(self, filepath, columns, rows, fmt='csv',
                 with_header=True, null_as_empty=True,
                 batch_size=DEFAULT_BATCH_SIZE, chunk_delay=0.0):
        super().__init__()
        self.setAutoDelete(False)
        self.filepath = filepath
        self.columns = list(columns or [])
        self.rows = list(rows or [])
        self.fmt = fmt
        self.with_header = bool(with_header)
        self.null_as_empty = bool(null_as_empty)
        self.batch_size = int(batch_size or DEFAULT_BATCH_SIZE)
        self.chunk_delay = float(chunk_delay or 0.0)
        self.signals = TransferSignals()
        self._cancel_event = threading.Event()

    def cancel(self):
        self._cancel_event.set()

    @property
    def cancel_requested(self) -> bool:
        return self._cancel_event.is_set()

    def run(self):
        try:
            written = write_export_chunked(
                self.filepath, self.columns, self.rows, fmt=self.fmt,
                with_header=self.with_header,
                null_as_empty=self.null_as_empty,
                batch_size=self.batch_size,
                progress_cb=lambda done, total:
                    self.signals.progress.emit(int(done), int(total)),
                cancel_cb=lambda: self._cancel_event.is_set(),
                chunk_delay=self.chunk_delay)
            self.signals.finished.emit({
                'ok': True, 'cancelled': False, 'written': written,
                'path': self.filepath, 'file_deleted': False})
        except ExportCancelled:
            self.signals.finished.emit({
                'ok': False, 'cancelled': True, 'written': 0,
                'path': self.filepath, 'file_deleted': True})
        except BaseException as exc:                  # noqa: BLE001
            _safe_remove(self.filepath)
            self.signals.failed.emit(str(exc))


# ==================== 导入对话框 ====================

class ImportCSVDialog(QDialog):
    """CSV 导入到表（P1-4：分块后台导入 + 进度 + 取消 + 失败报告）"""

    def __init__(self, conn, table_name, columns, parent=None):
        super().__init__(parent)
        self.conn = conn
        self.table_name = table_name
        self.columns = columns
        self.csv_path = ""
        self.headers = []
        self.rows = []
        self.last_import_result = None
        self._worker = None
        self._signals = None
        self._progress = None
        self._pool = QThreadPool(self)
        self._pool.setMaxThreadCount(1)
        self.batch_size = DEFAULT_BATCH_SIZE   # 可被测试/调用方覆盖
        self.setWindowTitle(f"导入 CSV → {table_name}")
        self.resize(760, 560)
        self._build_ui()

    def _build_ui(self):
        layout = QVBoxLayout(self)

        # 文件行
        file_row = QHBoxLayout()
        file_row.addWidget(QLabel("CSV 文件:"))
        self.file_edit = QLineEdit()
        self.file_edit.setReadOnly(True)
        file_row.addWidget(self.file_edit)
        browse_btn = QPushButton("浏览...")
        browse_btn.clicked.connect(self._browse)
        file_row.addWidget(browse_btn)
        layout.addLayout(file_row)

        # 选项
        opt_row = QHBoxLayout()
        self.skip_header = QCheckBox("首行是表头")
        self.skip_header.setChecked(True)
        self.skip_header.stateChanged.connect(self._reload_preview)
        opt_row.addWidget(self.skip_header)

        self.truncate = QCheckBox("导入前清空表")
        opt_row.addWidget(self.truncate)
        opt_row.addStretch()
        layout.addLayout(opt_row)

        # 列映射信息
        self.map_label = QLabel("")
        self.map_label.setStyleSheet("color: #969696;")
        layout.addWidget(self.map_label)

        # 预览
        layout.addWidget(QLabel("预览（前 20 行）:"))
        self.preview = QTableWidget()
        layout.addWidget(self.preview)

        # 失败报告按钮（导入后才可能启用）
        report_row = QHBoxLayout()
        self.failure_btn = QPushButton("导出失败行 CSV...")
        self.failure_btn.setEnabled(False)
        self.failure_btn.clicked.connect(self._export_failures)
        report_row.addWidget(self.failure_btn)
        report_row.addStretch()
        layout.addLayout(report_row)

        # 按钮
        buttons = QDialogButtonBox(
            QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        buttons.button(QDialogButtonBox.Ok).setText("开始导入")
        buttons.button(QDialogButtonBox.Cancel).setText("取消")
        buttons.accepted.connect(self._do_import)
        buttons.rejected.connect(self.reject)
        layout.addWidget(buttons)

    def _browse(self):
        path, _ = QFileDialog.getOpenFileName(
            self, "选择 CSV 文件", "",
            "CSV 文件 (*.csv);;所有文件 (*.*)")
        if not path:
            return
        self.csv_path = path
        self.file_edit.setText(path)
        self._load_csv()
        self._reload_preview()

    def _load_csv(self):
        self.headers = []
        self.rows = []
        if not self.csv_path:
            return
        try:
            try:
                with open(self.csv_path, 'r', encoding='utf-8', newline='') as f:
                    all_rows = list(csv.reader(f))
            except UnicodeDecodeError:
                with open(self.csv_path, 'r', encoding='gbk', newline='') as f:
                    all_rows = list(csv.reader(f))

            if not all_rows:
                return

            if self.skip_header.isChecked():
                self.headers = all_rows[0]
                self.rows = all_rows[1:]
            else:
                self.headers = [f"col{i+1}" for i in range(len(all_rows[0]))]
                self.rows = all_rows
        except Exception as e:
            QMessageBox.critical(self, "读取失败", str(e))

    def _reload_preview(self):
        if not self.rows:
            self.preview.clear()
            self.map_label.setText("")
            return

        n = min(20, len(self.rows))
        cols = len(self.rows[0]) if self.rows else 0
        self.preview.setColumnCount(cols)
        self.preview.setHorizontalHeaderLabels(
            self.headers or [f"col{i+1}" for i in range(cols)])
        self.preview.setRowCount(n)
        for ri in range(n):
            for ci in range(cols):
                v = self.rows[ri][ci] if ci < len(self.rows[ri]) else ''
                self.preview.setItem(ri, ci, QTableWidgetItem(str(v)))
        self.preview.resizeColumnsToContents()

        self.map_label.setText(
            f"CSV 有 {cols} 列 | 表有 {len(self.columns)} 列 | "
            f"共 {len(self.rows)} 行待导入")

    def build_import_sql(self):
        n_cols = len(self.columns)
        cols_str = ', '.join(f"`{c}`" for c in self.columns)
        placeholders = ', '.join(['%s'] * n_cols)
        return (f"INSERT INTO `{self.table_name}` ({cols_str}) "
                f"VALUES ({placeholders})")

    def _do_import(self):
        if not self.rows:
            QMessageBox.warning(self, "提示", "没有可导入的数据")
            return
        if self._worker is not None:
            return

        n_cols = len(self.columns)
        csv_cols = len(self.rows[0]) if self.rows else 0
        if csv_cols != n_cols:
            ret = QMessageBox.question(
                self, "列数不符",
                f"CSV 有 {csv_cols} 列，表有 {n_cols} 列。\n"
                f"将按顺序映射，多余列忽略，缺失列填 NULL。是否继续？")
            if ret != QMessageBox.Yes:
                return

        if self.truncate.isChecked():
            ret = QMessageBox.question(
                self, "确认", f"确定要清空表 {self.table_name} 吗？")
            if ret != QMessageBox.Yes:
                return
            _, err = self.conn.execute(f"TRUNCATE TABLE `{self.table_name}`")
            if err:
                QMessageBox.critical(self, "错误", err)
                return

        sql = self.build_import_sql()
        params_list = normalize_import_params(self.rows, n_cols)
        self.last_import_result = None
        self.failure_btn.setEnabled(False)
        self._progress = QProgressDialog("导入中...", "取消", 0,
                                         len(params_list), self)
        self._progress.setWindowModality(Qt.WindowModal)
        self._progress.setMinimumDuration(0)
        self._progress.setValue(0)

        worker = ImportWorker(self.conn, sql, params_list,
                              batch_size=self.batch_size)
        self._worker = worker
        self._signals = worker.signals
        worker.signals.progress.connect(self._on_import_progress)
        worker.signals.finished.connect(self._on_import_finished)
        worker.signals.failed.connect(self._on_import_failed)
        self._progress.canceled.connect(worker.cancel)
        self._pool.start(worker)

    def _on_import_progress(self, done, total):
        if self._progress is not None:
            self._progress.setMaximum(int(total) or 1)
            self._progress.setValue(int(done))

    def _on_import_finished(self, result):
        self.last_import_result = dict(result or {})
        self._close_progress()
        self._worker = None
        failures = self.last_import_result.get('failures') or []
        success = self.last_import_result.get('success', 0)
        cancelled = self.last_import_result.get('cancelled', False)
        self.failure_btn.setEnabled(bool(failures))
        msg = f"成功 {success} 行，失败 {len(failures)} 行"
        if cancelled:
            msg += (f"\n\n已取消：已导入 {success} 行（**已提交批次不回滚**），"
                    f"后续批次未执行。")
        if failures:
            msg += "\n\n失败行可通过「导出失败行 CSV...」保存（含行号/原因）。"
        if cancelled:
            QMessageBox.information(self, "导入已取消", msg)
        else:
            QMessageBox.information(self, "完成", msg)
            if not failures:
                self.accept()

    def _on_import_failed(self, error):
        self._close_progress()
        self._worker = None
        QMessageBox.critical(self, "导入失败", str(error))

    def _close_progress(self):
        if self._progress is not None:
            try:
                self._progress.close()
            except Exception:
                pass
            self._progress = None

    def _export_failures(self, filepath=None):
        failures = (self.last_import_result or {}).get('failures') or []
        if not failures:
            QMessageBox.information(self, "提示", "本次导入没有失败行")
            return False
        if not filepath:
            filepath, _ = QFileDialog.getSaveFileName(
                self, "导出失败行", f"{self.table_name}_import_failures.csv",
                "CSV 文件 (*.csv)")
        if not filepath:
            return False
        count = write_failure_report(failures, filepath)
        QMessageBox.information(
            self, "完成", f"已导出 {count} 条失败行到:\n{filepath}")
        return True


# ==================== 导出对话框 ====================

class ExportCSVDialog(QDialog):
    """导出结果/表到 CSV 或 XLSX（P1-4：后台分块 + 进度 + 取消）"""

    def __init__(self, columns, rows, default_name="export.csv", parent=None):
        super().__init__(parent)
        self.columns = columns
        self.rows = rows
        self._worker = None
        self._signals = None
        self._progress = None
        self._pool = QThreadPool(self)
        self._pool.setMaxThreadCount(1)
        self.setWindowTitle("导出结果")
        self.resize(560, 300)

        layout = QVBoxLayout(self)
        layout.addWidget(QLabel(
            f"共 {len(rows)} 行，{len(columns)} 列"))

        fmt_row = QHBoxLayout()
        fmt_row.addWidget(QLabel("格式:"))
        self.fmt_combo = QComboBox()
        self.fmt_combo.addItem("CSV", 'csv')
        self.fmt_combo.addItem("Excel (XLSX)", 'xlsx')
        self.fmt_combo.currentIndexChanged.connect(self._on_format_changed)
        fmt_row.addWidget(self.fmt_combo)
        fmt_row.addStretch()
        layout.addLayout(fmt_row)

        row = QHBoxLayout()
        row.addWidget(QLabel("保存到:"))
        self.path_edit = QLineEdit(default_name)
        row.addWidget(self.path_edit)
        browse_btn = QPushButton("浏览...")
        browse_btn.clicked.connect(self._browse)
        row.addWidget(browse_btn)
        layout.addLayout(row)

        self.with_header = QCheckBox("包含表头")
        self.with_header.setChecked(True)
        layout.addWidget(self.with_header)

        self.null_as_empty = QCheckBox("NULL 导出为空字符串")
        self.null_as_empty.setChecked(True)
        layout.addWidget(self.null_as_empty)

        layout.addStretch()

        buttons = QDialogButtonBox(
            QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        buttons.button(QDialogButtonBox.Ok).setText("导出")
        buttons.button(QDialogButtonBox.Cancel).setText("取消")
        buttons.accepted.connect(self._do_export)
        buttons.rejected.connect(self.reject)
        layout.addWidget(buttons)

    def export_format(self) -> str:
        return self.fmt_combo.currentData() or 'csv'

    def _on_format_changed(self, _index):
        fmt = self.export_format()
        path = self.path_edit.text().strip()
        if path:
            base = os.path.splitext(path)[0]
            self.path_edit.setText(base + ('.xlsx' if fmt == 'xlsx' else '.csv'))

    def _browse(self):
        fmt = self.export_format()
        if fmt == 'xlsx':
            path, _ = QFileDialog.getSaveFileName(
                self, "保存 Excel", self.path_edit.text(),
                "Excel 文件 (*.xlsx)")
        else:
            path, _ = QFileDialog.getSaveFileName(
                self, "保存 CSV", self.path_edit.text(),
                "CSV 文件 (*.csv)")
        if path:
            self.path_edit.setText(path)

    def _do_export(self):
        path = self.path_edit.text().strip()
        if not path:
            QMessageBox.warning(self, "提示", "请填写保存路径")
            return
        if self._worker is not None:
            return
        fmt = self.export_format()
        want_ext = '.xlsx' if fmt == 'xlsx' else '.csv'
        if not path.lower().endswith(want_ext):
            path += want_ext
            self.path_edit.setText(path)

        self._progress = QProgressDialog("导出中...", "取消", 0,
                                         max(1, len(self.rows)), self)
        self._progress.setWindowModality(Qt.WindowModal)
        self._progress.setMinimumDuration(0)
        self._progress.setValue(0)

        worker = ExportWorker(
            path, self.columns, self.rows, fmt=fmt,
            with_header=self.with_header.isChecked(),
            null_as_empty=self.null_as_empty.isChecked(),
            batch_size=DEFAULT_BATCH_SIZE)
        self._worker = worker
        self._signals = worker.signals
        worker.signals.progress.connect(self._on_export_progress)
        worker.signals.finished.connect(self._on_export_finished)
        worker.signals.failed.connect(self._on_export_failed)
        self._progress.canceled.connect(worker.cancel)
        self._pool.start(worker)

    def _on_export_progress(self, done, total):
        if self._progress is not None:
            self._progress.setMaximum(int(total) or 1)
            self._progress.setValue(int(done))

    def _on_export_finished(self, result):
        self._close_progress()
        self._worker = None
        result = dict(result or {})
        if result.get('cancelled'):
            QMessageBox.information(
                self, "已取消",
                "导出已取消，不完整文件已删除：\n"
                f"{result.get('path')}")
            return
        QMessageBox.information(
            self, "完成",
            f"已导出 {result.get('written', 0)} 行到:\n"
            f"{result.get('path')}")
        self.accept()

    def _on_export_failed(self, error):
        self._close_progress()
        self._worker = None
        QMessageBox.critical(
            self, "失败",
            f"导出失败（不完整文件已删除）：\n{error}")

    def _close_progress(self):
        if self._progress is not None:
            try:
                self._progress.close()
            except Exception:
                pass
            self._progress = None
