# -*- coding: utf-8 -*-
"""P1-4 导入/导出进度与取消 + P1-5 DataEditor 类型编辑/原子提交。

offscreen Qt + 假连接（记录 SQL/事务调用，不依赖真实服务端）。覆盖：
1. 导入分块（1000 行 → 2×500 批）、进度递增、取消后不再执行后续批次且
   显式返回“已取消/已导入 N 行（不回滚）”、批失败逐行重试收集失败行、
   失败报告 CSV 内容；
2. 导出 CSV/XLSX 分块、进度、取消删除不完整文件；ExportWorker 取消信号；
3. DataEditor 列编辑器类型（BOOL→下拉、DATE/DATETIME→日期控件、数值→校验
   输入框、JSON→文本框）；非法 JSON/日期/数值显式报错且**不提交**；
   NULL 显式设置；
4. 原子提交：BEGIN → 逐条 UPDATE → COMMIT；任一失败 ROLLBACK；
   单条修改同样走事务路径。
"""
import json
import os
import sys
import time

os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')
sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..'))

import pytest

_APP = None


def qapp():
    global _APP
    from PySide6.QtWidgets import QApplication
    app = QApplication.instance()
    if app is None:
        app = QApplication([])
    _APP = app
    return _APP


def pump(seconds=0.03):
    app = qapp()
    end = time.time() + seconds
    while time.time() < end:
        app.processEvents()
        time.sleep(0.005)


def wait_until(pred, timeout=8.0):
    app = qapp()
    end = time.time() + timeout
    while time.time() < end:
        app.processEvents()
        if pred():
            return True
        time.sleep(0.01)
    app.processEvents()
    return bool(pred())


# ==================== 假连接 ====================

class FakeImportConn:
    """execute_many 分块记录；可注入整批失败 / 逐行失败。"""

    def __init__(self, fail_batches=None, fail_rows=None, delay=0.0):
        self.execute_many_calls = []
        self.execute_calls = []
        self.fail_batches = set(fail_batches or ())
        self.fail_rows = set(fail_rows or ())
        self.delay = delay
        self._row_seq = 0

    def execute_many(self, sql, batch):
        if self.delay:
            time.sleep(self.delay)
        self.execute_many_calls.append(list(batch))
        idx = len(self.execute_many_calls) - 1
        if idx in self.fail_batches:
            return 0, f'batch {idx} failed'
        return len(batch), None

    def execute(self, sql, params=None):
        up = sql.strip().upper()
        if up in ('BEGIN', 'COMMIT', 'ROLLBACK'):
            self.execute_calls.append(up)
            return 0, None
        self._row_seq += 1
        self.execute_calls.append((sql, params))
        if self._row_seq in self.fail_rows:
            return 0, f'row {self._row_seq} failed'
        return 1, None


COLUMNS_META = [
    {'name': 'id', 'type': 'INT', 'null': 'NO', 'key': 'PRI',
     'default': None, 'extra': ''},
    {'name': 'name', 'type': 'VARCHAR(32)', 'null': 'YES', 'key': '',
     'default': None, 'extra': ''},
    {'name': 'born', 'type': 'DATE', 'null': 'YES', 'key': '',
     'default': None, 'extra': ''},
    {'name': 'active', 'type': 'BOOL', 'null': 'YES', 'key': '',
     'default': None, 'extra': ''},
    {'name': 'score', 'type': 'DOUBLE', 'null': 'YES', 'key': '',
     'default': None, 'extra': ''},
    {'name': 'payload', 'type': 'JSON', 'null': 'YES', 'key': '',
     'default': None, 'extra': ''},
]


class FakeEditorConn:
    """DataEditor 用：query 返回元数据/数据，execute 记录事务序列。"""

    def __init__(self, fail_at=None, rows=None):
        self.executed = []
        self.fail_at = fail_at
        self._stmt_no = 0
        self.rows = rows or [[1, 'alice', '2026-01-02', 1, 9.5,
                              '{"a": 1}']]

    def query(self, sql, params=None):
        up = sql.strip().upper()
        if 'COUNT(*)' in up:
            return ['COUNT(*)'], [[len(self.rows)]]
        if up.startswith('SELECT'):
            return [c['name'] for c in COLUMNS_META], \
                [tuple(r) for r in self.rows]
        return [], []

    def execute(self, sql, params=None):
        up = sql.strip().upper()
        self.executed.append(sql)
        if up in ('BEGIN', 'COMMIT', 'ROLLBACK'):
            return 0, None
        self._stmt_no += 1
        if self.fail_at is not None and self._stmt_no >= self.fail_at:
            return 0, f'DB error at stmt {self._stmt_no}'
        return 1, None


def make_editor(fail_at=None):
    from lightdb_studio.ui.data_editor import DataEditor
    qapp()
    conn = FakeEditorConn(fail_at=fail_at)
    return DataEditor(conn, 'people', COLUMNS_META), conn


class _MsgBox:
    """替换 QMessageBox 静态方法：记录调用，避免模态阻塞。"""

    def __init__(self):
        self.calls = []

    def _rec(self, kind):
        def fn(*args, **kwargs):
            self.calls.append({'kind': kind,
                               'text': ' | '.join(str(a) for a in args)})
            from PySide6.QtWidgets import QMessageBox
            return QMessageBox.Yes
        return fn

    def texts(self, kind=None):
        return [c['text'] for c in self.calls
                if kind is None or c['kind'] == kind]


@pytest.fixture
def msgbox(monkeypatch):
    box = _MsgBox()
    from PySide6.QtWidgets import QMessageBox
    monkeypatch.setattr(QMessageBox, 'information',
                        staticmethod(box._rec('information')))
    monkeypatch.setattr(QMessageBox, 'warning',
                        staticmethod(box._rec('warning')))
    monkeypatch.setattr(QMessageBox, 'critical',
                        staticmethod(box._rec('critical')))
    monkeypatch.setattr(QMessageBox, 'question',
                        staticmethod(box._rec('question')))
    return box


# ==================== 1. 导入分块 / 进度 / 取消 / 失败行 ====================

def test_import_chunked_1000_rows():
    from lightdb_studio.ui.import_export import (
        execute_chunked_import, DEFAULT_BATCH_SIZE)
    conn = FakeImportConn()
    params = [(i, f'name{i}') for i in range(1000)]
    progress = []
    result = execute_chunked_import(
        conn, 'INSERT INTO t (a, b) VALUES (%s, %s)', params,
        batch_size=DEFAULT_BATCH_SIZE,
        progress_cb=lambda d, t: progress.append((d, t)))
    print(f"  → result={result['success']}/{result['total']} 批次数="
          f"{len(conn.execute_many_calls)} 进度={progress[-3:]}")
    assert DEFAULT_BATCH_SIZE == 500
    assert len(conn.execute_many_calls) == 2
    assert [len(b) for b in conn.execute_many_calls] == [500, 500]
    assert result['success'] == 1000 and result['total'] == 1000
    assert result['failures'] == [] and result['cancelled'] is False
    assert progress[0] == (500, 1000) and progress[-1] == (1000, 1000)
    assert progress == sorted(progress)


def test_import_cancel_stops_later_batches():
    from lightdb_studio.ui.import_export import execute_chunked_import
    conn = FakeImportConn()
    params = [(i,) for i in range(1000)]
    done = {'n': 0}
    result = execute_chunked_import(
        conn, 'INSERT', params, batch_size=500,
        progress_cb=lambda d, t: done.__setitem__('n', d),
        cancel_cb=lambda: done['n'] >= 500)
    print(f"  → cancelled={result['cancelled']} success={result['success']} "
          f"批次数={len(conn.execute_many_calls)}")
    assert result['cancelled'] is True
    assert result['success'] == 500, '取消前已提交批次应保留'
    assert len(conn.execute_many_calls) == 1, '取消后不应执行后续批次'


def test_import_failure_rows_collected_and_report(tmp_path):
    from lightdb_studio.ui.import_export import (
        execute_chunked_import, write_failure_report)
    conn = FakeImportConn(fail_batches={0}, fail_rows={3, 5})
    params = [(i,) for i in range(1000)]
    result = execute_chunked_import(conn, 'INSERT', params, batch_size=500)
    print(f"  → success={result['success']} failures="
          f"{[(f['row'], f['error']) for f in result['failures']]}")
    assert result['cancelled'] is False
    assert len(result['failures']) == 2
    assert [f['row'] for f in result['failures']] == [3, 5]
    assert result['success'] == 998
    report = tmp_path / 'failures.csv'
    count = write_failure_report(result['failures'], report)
    assert count == 2 and report.exists()
    text = report.read_text(encoding='utf-8-sig')
    assert '行号' in text and '错误' in text
    assert '3' in text and 'row 3 failed' in text


def test_import_worker_signals_in_background():
    from PySide6.QtCore import QThreadPool
    from lightdb_studio.ui.import_export import ImportWorker
    qapp()
    conn = FakeImportConn()
    worker = ImportWorker(conn, 'INSERT', [(i,) for i in range(1200)],
                          batch_size=500)
    progresses = []
    finished = []
    worker.signals.progress.connect(lambda d, t: progresses.append((d, t)))
    worker.signals.finished.connect(finished.append)
    QThreadPool.globalInstance().start(worker)
    assert wait_until(lambda: bool(finished), 8.0), 'worker 未完成'
    print(f"  → 进度信号 {progresses}; finished={finished[0]['success']}")
    assert finished[0]['success'] == 1200
    assert progresses and progresses[-1] == (1200, 1200)


def test_import_dialog_chunked_with_fake_conn(msgbox):
    from lightdb_studio.ui.import_export import ImportCSVDialog
    qapp()
    conn = FakeImportConn()
    dlg = ImportCSVDialog(conn, 't', ['a', 'b'])
    dlg.rows = [[f'v{i}', i] for i in range(1000)]
    dlg._do_import()
    assert wait_until(lambda: dlg.last_import_result is not None, 8.0)
    result = dlg.last_import_result
    print(f"  → dialog 导入 result={result.get('success')}")
    assert result['success'] == 1000
    assert len(conn.execute_many_calls) == 2
    assert not dlg.failure_btn.isEnabled()
    dlg.close()


def test_import_dialog_cancel_semantics(msgbox):
    from lightdb_studio.ui.import_export import ImportCSVDialog
    qapp()
    conn = FakeImportConn(delay=0.02)      # 每批 20ms，保证取消时机确定
    dlg = ImportCSVDialog(conn, 't', ['a'])
    dlg.batch_size = 100                    # 2000 行 → 20 批，每批 20ms
    dlg.rows = [[f'v{i}'] for i in range(2000)]
    dlg._do_import()
    # 第一批完成后取消（进度信号驱动）
    assert wait_until(
        lambda: dlg._progress is not None and dlg._progress.value() >= 100,
        8.0)
    if dlg._worker is not None:
        dlg._worker.cancel()
    assert wait_until(lambda: dlg.last_import_result is not None, 8.0)
    result = dlg.last_import_result
    print(f"  → 取消结果 cancelled={result['cancelled']} "
          f"success={result['success']}")
    assert result['cancelled'] is True
    assert 100 <= result['success'] < 2000
    assert any('取消' in t for t in msgbox.texts('information')), \
        msgbox.texts('information')
    dlg.close()


# ==================== 2. 导出分块 / 取消删除 ====================

def test_export_csv_success_chunked(tmp_path):
    from lightdb_studio.ui.import_export import write_export_chunked
    rows = [[i, None if i % 3 == 0 else f'n{i}'] for i in range(1200)]
    path = tmp_path / 'out.csv'
    progress = []
    written = write_export_chunked(
        str(path), ['id', 'name'], rows, fmt='csv', batch_size=500,
        progress_cb=lambda d, t: progress.append((d, t)))
    assert written == 1200
    lines = path.read_text(encoding='utf-8-sig').strip().split('\n')
    assert len(lines) == 1201
    assert lines[0] == 'id,name'
    assert progress[0] == (500, 1200) and progress[-1] == (1200, 1200)


def test_export_cancel_deletes_incomplete_file(tmp_path):
    from lightdb_studio.ui.import_export import (
        write_export_chunked, ExportCancelled)
    rows = [[i] for i in range(1000)]
    path = tmp_path / 'cancel.csv'
    done = {'n': 0}

    def progress(d, t):
        done['n'] = d

    with pytest.raises(ExportCancelled):
        write_export_chunked(
            str(path), ['id'], rows, fmt='csv', batch_size=500,
            progress_cb=progress, cancel_cb=lambda: done['n'] >= 500)
    print(f"  → 取消后文件存在={path.exists()}")
    assert not path.exists(), '取消导出后必须删除不完整文件'


def test_export_xlsx_success_and_cancel(tmp_path):
    from lightdb_studio.ui.import_export import (
        write_export_chunked, ExportCancelled)
    rows = [[i, f'n{i}'] for i in range(120)]
    path = tmp_path / 'out.xlsx'
    write_export_chunked(str(path), ['id', 'name'], rows, fmt='xlsx',
                         batch_size=50)
    import openpyxl
    wb = openpyxl.load_workbook(str(path))
    ws = wb.active
    assert ws.max_row == 121 and ws.max_column == 2
    cancel_path = tmp_path / 'cancel.xlsx'
    with pytest.raises(ExportCancelled):
        write_export_chunked(str(cancel_path), ['id'], rows, fmt='xlsx',
                             batch_size=10, cancel_cb=lambda: True)
    assert not cancel_path.exists()


def test_export_worker_cancel_signal(tmp_path):
    from PySide6.QtCore import QThreadPool
    from lightdb_studio.ui.import_export import ExportWorker
    from PySide6.QtWidgets import QApplication
    qapp()
    path = tmp_path / 'worker_cancel.csv'
    rows = [[i] for i in range(2000)]
    worker = ExportWorker(str(path), ['id'], rows, fmt='csv',
                          batch_size=100, chunk_delay=0.01)
    progresses = []
    finished = []
    worker.signals.progress.connect(lambda d, t: progresses.append(d))
    worker.signals.finished.connect(finished.append)
    QThreadPool.globalInstance().start(worker)
    assert wait_until(lambda: bool(progresses), 5.0)
    worker.cancel()
    assert wait_until(lambda: bool(finished), 8.0)
    print(f"  → ExportWorker 取消结果={finished[0]} 文件存在={path.exists()}")
    assert finished[0]['cancelled'] is True
    assert finished[0]['file_deleted'] is True
    assert not path.exists()


# ==================== 3. DataEditor 类型编辑/校验 ====================

def test_editor_kind_mapping_and_delegate_widgets():
    from PySide6.QtWidgets import QComboBox, QDateEdit, QDateTimeEdit, \
        QLineEdit, QStyleOptionViewItem
    from lightdb_studio.ui.data_editor import (
        editor_kind_for_type, TypedCellDelegate)
    assert editor_kind_for_type('INT') == 'int'
    assert editor_kind_for_type('BIGINT') == 'int'
    assert editor_kind_for_type('DOUBLE') == 'float'
    assert editor_kind_for_type('DATE') == 'date'
    assert editor_kind_for_type('DATETIME') == 'datetime'
    assert editor_kind_for_type('BOOL') == 'bool'
    assert editor_kind_for_type('JSON') == 'json'
    assert editor_kind_for_type('VARCHAR(20)') == 'text'

    editor, _conn = make_editor()
    qapp()
    delegate = TypedCellDelegate(COLUMNS_META, editor.table)
    model = editor.table.model()
    classes = {}
    for ci, name in enumerate(['id', 'name', 'born', 'active', 'score',
                               'payload']):
        w = delegate.createEditor(editor.table, QStyleOptionViewItem(),
                                  model.index(0, ci))
        classes[name] = type(w).__name__
        w.deleteLater()
    print(f"  → 列编辑器类型={classes}")
    assert classes['active'] == 'QComboBox'
    assert classes['born'] == 'QDateEdit'
    assert classes['id'] == 'QLineEdit'
    assert classes['score'] == 'QLineEdit'
    assert classes['payload'] == 'QLineEdit'
    editor.close()


def test_validate_typed_value_rules():
    from lightdb_studio.ui.data_editor import validate_typed_value as v
    assert v('INT', '12') == (True, 12, '')
    ok, _val, err = v('INT', '12.5')
    assert not ok and '非法' in err
    assert v('DOUBLE', '-3.5')[0] is True
    ok, val, err = v('DATE', '2026/10/21')
    assert ok and val == '2026-10-21'
    ok, _val, err = v('DATE', '2026-13-40')
    assert not ok and '日期格式' in err
    ok, val, _e = v('DATETIME', '2026-10-21 09:30:00')
    assert ok and val == '2026-10-21 09:30:00'
    assert v('BOOL', 'TRUE') == (True, True, '')
    assert v('BOOL', 'FALSE') == (True, False, '')
    ok, _val, err = v('BOOL', 'maybe')
    assert not ok
    ok, val, _e = v('JSON', '{"a": 1}')
    assert ok and val == '{"a": 1}'
    ok, _val, err = v('JSON', '{bad json')
    assert not ok and 'JSON 非法' in err
    assert v('JSON', 'NULL') == (True, None, '')
    assert v('VARCHAR(10)', 'hello') == (True, 'hello', '')


def test_invalid_json_blocks_commit(msgbox):
    from PySide6.QtWidgets import QMessageBox
    editor, conn = make_editor()
    qapp()
    before = len(conn.executed)
    # payload 列（index 5）改写为非法 JSON
    editor.table.item(0, 5).setText('{not-json')
    qapp().processEvents()
    assert editor.invalid_cells, '非法 JSON 未进入 invalid_cells'
    assert any('非法值' in t for t in
               [editor.pending_label.text()]) or '非法值' in \
        editor.pending_label.text()
    editor._save()
    qapp().processEvents()
    assert '存在非法值' in ' '.join(msgbox.texts('critical')), \
        msgbox.texts('critical')
    assert len(conn.executed) == before, '存在非法值时不应执行任何 SQL'
    editor.close()


def test_set_cell_null_explicit():
    editor, _conn = make_editor()
    qapp()
    assert editor.set_cell_null(0, 1) is True
    assert editor.table.item(0, 1).text() == 'NULL'
    assert editor._row_data[0]['values'][1] is None
    assert 'name' in editor.modified.get(0, {}), editor.modified
    editor.close()


# ==================== 4. 原子提交 ====================

def test_execute_atomic_commit_sequence():
    from lightdb_studio.ui.data_editor import execute_atomic
    conn = FakeEditorConn()
    result = execute_atomic(conn, ['UPDATE a SET x=1',
                                   'UPDATE b SET y=2'])
    print(f"  → sequence={conn.executed} result={result}")
    assert conn.executed == ['BEGIN', 'UPDATE a SET x=1',
                             'UPDATE b SET y=2', 'COMMIT']
    assert result['ok'] is True and result['executed'] == 2
    assert result['rolled_back'] is False


def test_execute_atomic_rollback_on_failure():
    from lightdb_studio.ui.data_editor import execute_atomic
    conn = FakeEditorConn(fail_at=2)
    result = execute_atomic(conn, ['UPDATE a SET x=1',
                                   'UPDATE b SET y=2',
                                   'UPDATE c SET z=3'])
    print(f"  → failure sequence={conn.executed} result={result}")
    assert conn.executed[0] == 'BEGIN'
    assert conn.executed[-1] == 'ROLLBACK'
    assert result['ok'] is False and result['rolled_back'] is True
    assert result['executed'] == 1, '失败时已执行语句数应为 1'
    assert any('DB error' in e for e in result['errors'])


def test_dataeditor_commit_statements_atomic_success(msgbox):
    editor, conn = make_editor()
    qapp()
    editor.table.item(0, 1).setText('alice2')
    editor.table.item(0, 4).setText('12.5')
    qapp().processEvents()
    stmts = editor.build_pending_statements()
    assert len(stmts) == 1, stmts      # 同一行两个单元格合并成一条 UPDATE
    result = editor.commit_statements(stmts, ask=False)
    print(f"  → 提交序列={conn.executed}")
    assert result['ok'] is True
    assert conn.executed[0] == 'BEGIN'
    assert conn.executed[-1] == 'COMMIT'
    assert any(s.startswith('UPDATE `people`') for s in conn.executed)
    assert '成功' in ' '.join(msgbox.texts('information'))
    editor.close()


def test_dataeditor_commit_rollback_path(msgbox):
    editor, conn = make_editor(fail_at=1)
    qapp()
    editor.table.item(0, 1).setText('bob')
    qapp().processEvents()
    stmts = editor.build_pending_statements()
    result = editor.commit_statements(stmts, ask=False)
    print(f"  → 失败序列={conn.executed}")
    assert result['ok'] is False and result['rolled_back'] is True
    assert conn.executed[0] == 'BEGIN' and conn.executed[-1] == 'ROLLBACK'
    assert '已回滚' in ' '.join(msgbox.texts('critical'))
    editor.close()


def test_dataeditor_single_row_same_transaction_path(msgbox):
    """单条修改也必须 BEGIN/COMMIT（不允许逐条直写）。"""
    editor, conn = make_editor()
    qapp()
    editor.table.item(0, 2).setText('2026-02-03')
    qapp().processEvents()
    stmts = editor.build_pending_statements()
    assert len(stmts) == 1
    result = editor.commit_statements(stmts, ask=False)
    assert result['ok'] is True
    assert conn.executed[0] == 'BEGIN' and conn.executed[-1] == 'COMMIT'
    editor.close()


if __name__ == '__main__':
    print('本文件需 pytest（fixture）；直接运行仅做导入检查')
    sys.exit(0)
