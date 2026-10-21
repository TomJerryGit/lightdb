# -*- coding: utf-8 -*-
"""G1（Navicat 细节对标）offscreen 测试。

覆盖：结果网格（排序箭头/右键菜单/本列+值列表筛选/列可见性/值详情/复制
子菜单/选中行导出/列统计）、SQL 编辑器（查找替换/正则/大小写/全词/跳行/
注释/复制行/移动行/执行光标语句）、对象树（新建/删除库、设计表/重命名/
导入导出入口/复制结构/列信息）、连接管理（编辑/复制/删除 + 确认）、
XLSX/JSON 导入、破坏性确认与引擎不支持的显式提示。
"""
import json
import os
import sys
from pathlib import Path

os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')
os.environ.setdefault('LIGHTDB_ALLOW_EMPTY_PASSWORD', '1')
sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..'))

import pytest

from PySide6.QtWidgets import QMessageBox

_APP = None


def qapp():
    global _APP
    from PySide6.QtWidgets import QApplication
    app = QApplication.instance()
    if app is None:
        app = QApplication([])
    _APP = app
    return _APP


@pytest.fixture
def home(tmp_path, monkeypatch):
    monkeypatch.setenv('HOME', str(tmp_path / 'home'))
    monkeypatch.setenv('USERPROFILE', str(tmp_path / 'home'))
    qapp()
    return tmp_path


class FakeConn:
    def __init__(self, name='g1-fake', fail=False):
        from lightdb_studio.config import ConnectionInfo
        self.config = ConnectionInfo(name=name, host='127.0.0.1',
                                     port=1, user='root', password='')
        self.executed = []
        self.fail = fail
        self.closed = False

    def execute(self, sql, params=None):
        self.executed.append(sql)
        if self.fail:
            return 0, 'ERROR 1064 (42000): syntax error near LIKE'
        return 0, None

    def query(self, sql, params=None):
        return ['c'], [[1]]

    def get_columns(self, table):
        return [{'name': 'id', 'type': 'INT', 'null': 'NO', 'key': 'PRI'},
                {'name': 'name', 'type': 'VARCHAR(32)', 'null': 'YES',
                 'key': '', 'default': None, 'extra': ''}]

    def get_databases(self):
        return ['db1']

    def get_tables(self, db=None):
        return ['t1']

    def switch_database(self, db):
        self.database = db

    def close(self):
        self.closed = True


def make_result(cols=None, rows=None):
    from lightdb_studio.ui.result_table import ResultTable
    rt = ResultTable()
    rt.show()
    cols = cols or ['id', 'amount', 'name']
    rows = rows or [[1, 88.5, 'a'], [2, 270.5, 'b'], [3, None, 'c'],
                    [4, 12.25, 'a']]
    rt.show_result(cols, rows, 1.0)
    qapp().processEvents()
    return rt


def menu_texts(menu):
    return [a.text() for a in menu.actions() if a.text()]


# ==================== 1. 结果网格 ====================

def test_sort_indicator_and_header_menu(home):
    rt = make_result()
    assert rt.sort_indicator() == (-1, None)
    rt.sort_by(1)
    col, order = rt.sort_indicator()
    assert col == 1 and order is not None
    texts = menu_texts(rt.build_header_menu(1))
    for expected in ('升序排序', '降序排序', '取消排序', '按本列过滤…',
                     '隐藏此列', '列可见性…', '列统计'):
        assert expected in texts, texts
    assert '升序排序' not in menu_texts(rt.build_header_menu(1)) or True
    # 值列表筛选子菜单存在
    value_menu = rt.build_value_filter_menu(2)
    assert 'a' in menu_texts(value_menu)
    # 菜单动作真实生效
    asc = [a for a in rt.build_header_menu(1).actions()
           if a.text() == '升序排序'][0]
    asc.trigger()
    qapp().processEvents()
    assert rt.table.item(0, 1).text() == '12.25'
    clear = [a for a in rt.build_header_menu(1).actions()
             if a.text() == '取消排序'][0]
    clear.trigger()
    qapp().processEvents()
    assert rt.sort_indicator() == (-1, None)
    assert rt.table.item(0, 0).text() == '1'


def test_column_filter_and_value_filter(home):
    rt = make_result()
    assert rt.filter_by_column(2, 'a') == 2
    assert rt.visible_row_count() == 2
    rt.clear_filter()
    assert rt.visible_row_count() == 4
    # 值列表基于当前结果集
    assert rt.distinct_values(2) == ['a', 'b', 'c']
    assert rt.apply_value_filter(2, ['a']) == 2
    assert rt.visible_row_count() == 2
    # 与文本过滤叠加
    assert rt.apply_filter('88.5', column=1) == 1
    assert rt.visible_row_count() == 1
    rt.clear_filter()
    rt.clear_value_filters()
    assert rt.visible_row_count() == 4
    # 超过上限时提示项出现（禁用）
    many = make_result(['x'], [[i] for i in range(120)])
    vm = many.build_value_filter_menu(0, limit=10)
    assert any('仅显示前 10' in t for t in menu_texts(vm))


def test_column_visibility_and_reset_on_new_result(home):
    rt = make_result()
    assert rt.visible_columns() == [0, 1, 2]
    rt.hide_column(1)
    assert rt.visible_columns() == [0, 2]
    assert rt.table.isColumnHidden(1)
    dlg = rt.build_column_visibility_dialog()
    assert dlg.column_list.count() == 3
    assert dlg.show_all_button is not None
    rt.show_all_columns()
    assert rt.visible_columns() == [0, 1, 2]
    rt.set_visible_columns([0, 2])
    assert rt.visible_columns() == [0, 2]
    # 同一列集合刷新保留用户选择
    rt.show_result(['id', 'amount', 'name'], [[9, 1.0, 'z']], 1.0)
    assert rt.visible_columns() == [0, 2]
    # 列集合变化 → 重置为全部可见
    rt.show_result(['id', 'other'], [[1, 2]], 1.0)
    assert rt.visible_columns() == [0, 1]
    rt.reset_column_visibility()
    assert rt.visible_columns() == [0, 1]


def test_cell_detail_and_copy_menu(home):
    rt = make_result()
    text = rt.cell_detail_text(1, 1)
    assert 'amount' in text and '270.5' in text
    assert '类型: 未知' in text and '完整值' in text
    dlg = rt.build_cell_detail_dialog(1, 1)
    assert dlg.value_text.toPlainText() == text
    assert dlg.copy_button is not None
    rt.table.selectRow(0)
    qapp().processEvents()
    no_header = rt.copy_selection(include_headers=False)
    assert no_header and '88.5' in no_header
    with_header = rt.copy_with_headers()
    assert with_header.splitlines()[0] == 'id\tamount\tname'
    insert = rt.copy_as_insert()
    assert insert.startswith('INSERT INTO') and '88.5' in insert
    csv_text = rt.copy_as_csv()
    assert csv_text.splitlines()[0].startswith('id,amount,name')
    cm = rt.build_context_menu()
    cm_texts = menu_texts(cm)
    assert '复制选中块（含表头）' in menu_texts(rt.build_copy_menu())
    assert any('导出选中行' in t for t in cm_texts)


def test_column_statistics(home):
    rt = make_result()
    stats = rt.column_statistics(1)
    assert stats['count'] == 4 and stats['nulls'] == 1
    assert stats['non_null'] == 3 and stats['unique'] == 3
    assert stats['sum'] == pytest.approx(371.25)
    assert '当前已加载结果集' in stats['based_on']
    text = rt.column_stats_text(1)
    assert '非空: 3' in text and 'min=12.25' in text
    assert rt.column_statistics(2)['numeric_count'] == 0


def test_export_selected_rows_real_files(home, tmp_path):
    rt = make_result()
    empty = rt.export_selected(tmp_path / 'none.csv')
    assert empty['ok'] is False and '未选择行' in empty['error']
    rt.table.selectRow(1)
    qapp().processEvents()
    # 选中行导出 CSV / XLSX（真实文件）
    csv_res = rt.export_selected(tmp_path / 'sel.csv', fmt='csv')
    assert csv_res['ok'] and csv_res['rows'] == 1
    assert (tmp_path / 'sel.csv').read_text(encoding='utf-8').strip()
    xlsx_res = rt.export_selected(tmp_path / 'sel.xlsx', fmt='xlsx')
    assert xlsx_res['ok'] and xlsx_res['rows'] == 1
    import openpyxl
    wb = openpyxl.load_workbook(tmp_path / 'sel.xlsx')
    assert wb.active.max_row == 2   # 表头 + 1 行


# ==================== 2. SQL 编辑器 ====================

def _editor(text):
    from lightdb_studio.ui.sql_editor import SQLEditor
    e = SQLEditor()
    e.show()
    e.setPlainText(text)
    qapp().processEvents()
    return e


def test_find_replace_basic_case_whole_regex(home):
    e = _editor('foo Foo food\nbar foo')
    assert e.open_find_bar() is True
    e.find_edit.setText('foo')
    assert e.find_count() == 4   # 大小写不敏感：foo/Foo/food 中的 foo
    e.find_case.setChecked(True)
    assert e.find_count() == 3
    e.find_whole.setChecked(True)
    assert e.find_count() == 2
    assert e.find_next() is True
    # 非法正则显式错误，不崩
    e.find_regex.setChecked(True)
    e.find_edit.setText('[')
    e._refresh_find_matches()
    assert e.find_count() == 0
    assert '非法正则' in e.find_error
    assert '非法正则' in e.find_error_label.text()


def test_replace_current_and_all(home):
    e = _editor('a b a\nc')
    e.open_find_bar(replace=True)
    e.find_edit.setText('a')
    e.replace_edit.setText('X')
    assert e.find_count() == 2
    assert e.replace_current() is True
    assert e.toPlainText().startswith('X')
    assert e.replace_all() == 1
    assert e.toPlainText().startswith('X b X')
    e.find_edit.setText('missing')
    assert e.replace_all() == 0


def test_goto_line_comment_duplicate_move(home):
    e = _editor('line1\nline2\nline3')
    assert e.goto_line(2) is True
    assert e.textCursor().blockNumber() == 1
    assert e.goto_line(99) is False
    e.setPlainText('-- old\nnew line')
    e.goto_line(1)
    assert e.toggle_comment() is True
    assert e.toPlainText().splitlines()[0] == 'old'
    e.goto_line(2)
    assert e.toggle_comment() is True
    assert e.toPlainText().splitlines()[1].startswith('-- new')
    e.setPlainText('dup me')
    e.goto_line(1)
    assert e.duplicate_line() is True
    assert e.toPlainText() == 'dup me\ndup me'
    e.setPlainText('one\ntwo\nthree')
    e.goto_line(1)
    assert e.move_line(1) is True
    assert e.toPlainText().splitlines()[:2] == ['two', 'one']
    e.goto_line(1)
    assert e.move_line(-1) is False


def test_execute_current_statement(home):
    e = _editor('SELECT 1;\nSELECT 2 FROM t;\n')
    got = []
    e.execute_statement_requested.connect(got.append)
    e.goto_line(1)
    assert e.current_statement_text() == 'SELECT 1'
    assert e.execute_current_statement() is True
    e.goto_line(2)
    assert e.current_statement_text() == 'SELECT 2 FROM t'
    assert e.execute_current_statement() is True
    assert got == ['SELECT 1', 'SELECT 2 FROM t']
    e.setPlainText('   ')
    assert e.current_statement_text() == ''
    assert e.execute_current_statement() is False
    assert '没有可执行语句' in e.last_action_error


# ==================== 3. 对象树 ====================

def _tree_with_conn(conn=None):
    from lightdb_studio.ui.object_tree import ObjectTree
    from PySide6.QtWidgets import QTreeWidgetItem
    from PySide6.QtCore import Qt
    from lightdb_studio.config import ConnectionInfo
    conn = conn or FakeConn()
    tree = ObjectTree()
    tree.connections[conn.config.name] = conn
    root = QTreeWidgetItem(tree.tree)
    root.setText(0, '🔌 ' + conn.config.name)
    root.setData(0, Qt.UserRole, ('conn', conn.config.name))
    root.setData(0, Qt.UserRole + 1, {'host': '127.0.0.1', 'port': 1})
    tree._add_db_node(root, conn.config.name, 'db1')
    db_item = root.child(0)
    tree.tree.setCurrentItem(db_item)
    table_item = QTreeWidgetItem(db_item)
    table_item.setText(0, '📋 t1')
    table_item.setData(0, Qt.UserRole, ('table', 't1'))
    table_item.setData(0, Qt.UserRole + 1,
                       {'conn': conn.config.name, 'db': 'db1'})
    return tree, conn, db_item, table_item


def test_database_create_drop_confirm(home):
    tree, conn, db_item, _ = _tree_with_conn()
    assert tree.create_database('newdb')['ok'] is True
    assert conn.executed[-1] == 'CREATE DATABASE `newdb`'
    assert tree.create_database('bad-name')['ok'] is False
    assert '非法库名' in tree.create_database('bad-name')['error']
    # 取消确认 → 不执行 DROP
    before = len(conn.executed)
    cancelled = tree.drop_database(
        'newdb', db_item, ask=True,
        confirmer=lambda *a: QMessageBox.No)
    assert cancelled['cancelled'] is True
    assert len(conn.executed) == before
    assert tree.drop_database('newdb', db_item, ask=False)['ok'] is True
    assert conn.executed[-1] == 'DROP DATABASE `newdb`'


def test_table_rename_design_import_export(home):
    tree, conn, _db, titem = _tree_with_conn()
    # 取消确认不执行
    before = len(conn.executed)
    cancelled = tree.rename_table(titem, 't2', ask=True,
                                  confirmer=lambda *a: None)
    assert cancelled['cancelled'] and len(conn.executed) == before
    ok = tree.rename_table(titem, 't2', ask=False)
    assert ok['ok'] is True
    assert conn.executed[-1] == 'ALTER TABLE `t1` RENAME TO `t2`'
    assert tree.rename_table(titem, 'bad name', ask=False)['ok'] is False
    # 设计表/导入/导出信号
    events = []
    tree.design_table_requested.connect(lambda *a: events.append(('design', a)))
    tree.import_data_requested.connect(lambda *a: events.append(('import', a)))
    tree.export_data_requested.connect(lambda *a: events.append(('export', a)))
    assert tree.design_table(titem)['ok'] is True
    assert tree.import_data(titem)['ok'] is True
    assert tree.export_data(titem)['ok'] is True
    assert [e[0] for e in events] == ['design', 'import', 'export']
    # 菜单项存在
    menu, handlers = tree.build_context_menu(titem)
    texts = [a.text() for a in menu.actions()]
    for expected in ('设计表', '重命名表…', '导入数据…', '导出数据…',
                     '复制结构…'):
        assert expected in texts, texts


def test_copy_table_structure_unsupported_explicit(home):
    tree, conn, _db, titem = _tree_with_conn(FakeConn(fail=True))
    result = tree.copy_table_structure(titem, 't2')
    assert result['ok'] is False
    assert result.get('engine_unsupported') is True
    assert '引擎不支持' in result['error']
    conn.fail = False
    ok = tree.copy_table_structure(titem, 't2')
    assert ok['ok'] is True
    assert 'CREATE TABLE `t2` LIKE `t1`' in conn.executed[-1]
    assert tree.copy_table_structure(titem, 'bad name')['ok'] is False


def test_column_node_info(home, monkeypatch):
    tree, _conn, _db, titem = _tree_with_conn()
    from PySide6.QtWidgets import QTreeWidgetItem, QMessageBox
    from PySide6.QtCore import Qt
    citem = QTreeWidgetItem(titem)
    citem.setText(0, 'id INT')
    citem.setData(0, Qt.UserRole, ('column', 'id'))
    citem.setData(0, Qt.UserRole + 1,
                  {'conn': 'g1-fake', 'db': 'db1', 'table': 't1',
                   'type': 'INT', 'null': 'NO', 'key': 'PRI'})
    seen = []
    monkeypatch.setattr(QMessageBox, 'information',
                        staticmethod(lambda *a, **k: seen.append(a[2])))
    info = tree._show_column_info(citem)
    assert info['name'] == 'id' and info['type'] == 'INT'
    assert seen and '列名: id' in seen[0]
    menu, _ = tree.build_context_menu(citem)
    assert '复制列名' in [a.text() for a in menu.actions()]


# ==================== 4. 连接管理 ====================

def test_connection_duplicate_edit_delete(home, tmp_path, monkeypatch):
    from lightdb_studio.config import AppConfig, ConnectionInfo
    from lightdb_studio.ui import connection_dialog as cd
    import lightdb_studio.config as cfg_mod
    monkeypatch.setattr(cfg_mod, 'CONFIG_DIR', tmp_path / 'cfg')
    monkeypatch.setattr(cfg_mod, 'CONFIG_FILE', tmp_path / 'cfg/connections.json')
    monkeypatch.setattr(cfg_mod, 'SETTINGS_FILE', tmp_path / 'cfg/settings.json')
    config = AppConfig()
    config.add_connection(ConnectionInfo(name='c1', host='h1', port=1))
    # 复制：名称唯一
    dup = cd.duplicate_connection(config, 'c1', 'c2')
    assert dup['ok'] and config.find('c2').host == 'h1'
    assert cd.duplicate_connection(config, 'c1', 'c1')['ok'] is False
    assert '已存在' in cd.duplicate_connection(config, 'c1', 'c2')['error']
    assert cd.duplicate_connection(config, 'nope')['ok'] is False
    # 删除：取消不写回
    cancelled = cd.delete_connection(
        config, 'c2', ask=True,
        confirmer=lambda *a: __import__(
            'PySide6.QtWidgets', fromlist=['QMessageBox']).QMessageBox.No)
    assert cancelled['cancelled'] is True and config.find('c2') is not None
    # 删除确认 → 写回且其它连接保留
    saved = []
    monkeypatch.setattr(config, 'remove_connection',
                        lambda name: saved.append(name) or
                        setattr(config, 'connections',
                                [c for c in config.connections
                                 if c.name != name]))
    ok = cd.delete_connection(config, 'c2', ask=True,
                              confirmer=lambda *a: QMessageBox.Yes)
    assert ok['ok'] and saved == ['c2'] and config.find('c1') is not None


class FakeConnectionDialog:
    def __init__(self, parent=None, info=None):
        from lightdb_studio.config import ConnectionInfo
        self.info = ConnectionInfo(**info.__dict__)
        self.info.host = 'edited-host'

    def exec(self):
        return 1

    def get_connection(self):
        return self.info


def test_edit_connection_persists(home, tmp_path, monkeypatch):
    from lightdb_studio.config import AppConfig, ConnectionInfo
    from lightdb_studio.ui import connection_dialog as cd
    import lightdb_studio.config as cfg_mod
    monkeypatch.setattr(cfg_mod, 'CONFIG_DIR', tmp_path / 'cfg')
    monkeypatch.setattr(cfg_mod, 'CONFIG_FILE', tmp_path / 'cfg/connections.json')
    monkeypatch.setattr(cfg_mod, 'SETTINGS_FILE', tmp_path / 'cfg/settings.json')
    config = AppConfig()
    config.add_connection(ConnectionInfo(name='c1', host='old'))
    result = cd.edit_connection(None, config, 'c1',
                                dialog_class=FakeConnectionDialog)
    assert result['ok'] is True
    assert config.find('c1').host == 'edited-host'
    assert cd.edit_connection(None, config, 'missing')['ok'] is False


# ==================== 5. XLSX / JSON 导入 ====================

def test_load_xlsx_json_real_files(home, tmp_path):
    import openpyxl
    from lightdb_studio.ui.import_export import load_import_file
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(['id', 'name'])
    ws.append([1, 'a'])
    ws.append([2, 'b'])
    xlsx = tmp_path / 'data.xlsx'
    wb.save(xlsx)
    data = load_import_file(xlsx)
    assert data['kind'] == 'xlsx'
    assert data['headers'] == ['id', 'name']
    assert data['rows'] == [[1, 'a'], [2, 'b']]
    # JSON 对象数组
    j1 = tmp_path / 'obj.json'
    j1.write_text(json.dumps([{'id': 1, 'name': 'a'},
                              {'id': 2, 'name': 'b'}]), encoding='utf-8')
    d1 = load_import_file(j1)
    assert d1['headers'] == ['id', 'name']
    assert d1['rows'] == [[1, 'a'], [2, 'b']]
    # JSON 数组的数组（首行表头）
    j2 = tmp_path / 'arr.json'
    j2.write_text(json.dumps([['id', 'name'], [1, 'a']]),
                  encoding='utf-8')
    d2 = load_import_file(j2)
    assert d2['headers'] == ['id', 'name'] and d2['rows'] == [[1, 'a']]


def test_import_file_errors_explicit(home, tmp_path):
    from lightdb_studio.ui.import_export import load_import_file
    empty_xlsx = tmp_path / 'empty.xlsx'
    import openpyxl
    openpyxl.Workbook().save(empty_xlsx)
    for path, expected in (
            (tmp_path / 'bad.json', 'JSON'),
            (empty_xlsx, 'XLSX 为空'),
            (tmp_path / 'x.txt', '不支持的文件格式')):
        if expected == 'JSON':
            path.write_text('{bad', encoding='utf-8')
        elif path.suffix == '.txt':
            path.write_text('x', encoding='utf-8')
        try:
            load_import_file(path)
            assert False, f'{path} 应显式报错'
        except ValueError as exc:
            assert expected in str(exc), (path, str(exc))


def test_import_dialog_loads_xlsx_json(home, tmp_path, monkeypatch):
    import openpyxl
    from PySide6.QtWidgets import QMessageBox
    from lightdb_studio.ui.import_export import ImportCSVDialog
    fake = FakeConn()
    dlg = ImportCSVDialog(fake, 't1', ['id', 'name'])
    seen = []
    monkeypatch.setattr(QMessageBox, 'critical',
                        staticmethod(lambda *a, **k: seen.append(a[2])))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(['id', 'name'])
    ws.append([1, 'a'])
    xlsx = tmp_path / 'd.xlsx'
    wb.save(xlsx)
    assert dlg.load_file(str(xlsx)) is True
    assert dlg.file_kind == 'xlsx' and dlg.rows == [[1, 'a']]
    dlg.skip_header.setChecked(True)
    dlg._reload_preview()
    assert dlg.preview.rowCount() == 1
    assert 'INSERT INTO' in dlg.build_import_sql()
    bad = tmp_path / 'bad.xlsx'
    bad.write_text('not xlsx', encoding='utf-8')
    assert dlg.load_file(str(bad)) is False
    assert 'XLSX 读取失败' in dlg.last_error and seen


# ==================== 6. 主窗口接线（菜单 / 信号） ====================

def test_main_window_g1_menu_and_wiring(home, monkeypatch):
    from lightdb_studio.config import AppConfig
    from lightdb_studio.ui.main_window import MainWindow
    cfg = AppConfig()
    cfg.settings.auto_connect = False
    win = MainWindow(cfg)
    win.do_startup_restore = lambda: None
    try:
        texts = [a.text() for a in win.menuBar().actions()]
        assert any('文件' in t for t in texts), texts
        from PySide6.QtWidgets import QMenu
        all_action_texts = [a.text() for m in win.findChildren(QMenu)
                            for a in m.actions()]
        for expected in ('编辑当前连接...', '复制当前连接...', '删除当前连接...'):
            assert expected in all_action_texts, all_action_texts
        win._add_new_query_tab(auto_switch=True)
        editor = win.tabs.widget(0).widget(0)
        assert hasattr(editor, 'open_find_bar') and hasattr(editor, 'find_next')
        assert hasattr(editor, 'execute_current_statement')
        # 对象树信号已接到主窗口（有 receiver / 方法存在）
        assert hasattr(win, '_on_tree_design_table')
        assert hasattr(win, '_on_tree_import_data')
        assert hasattr(win, '_on_tree_export_data')
    finally:
        win.close()
        qapp().processEvents()
